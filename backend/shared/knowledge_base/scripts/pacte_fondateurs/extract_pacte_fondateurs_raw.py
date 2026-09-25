"""Extraction brute (Phase 0 — préparation), service 1.1.13 pacte de fondateurs
(Groupe 5) / réutilisable tel quel par 1.1.4 (statuts).

Lancement (une fois par fichier source), depuis `wink/` :

    python -m backend.shared.knowledge_base.scripts.pacte_fondateurs.extract_pacte_fondateurs_raw \
        --source "<chemin vers le .docx>" --label modele_principal

    python -m backend.shared.knowledge_base.scripts.pacte_fondateurs.extract_pacte_fondateurs_raw \
        --source "<chemin vers la notice explicative .docx>" --label notice_explicative

Écrit chaque source lue dans `backend/shared/knowledge_base/scripts/data/
pacte_fondateurs/raw_<label>.json` — un paragraphe du `.docx` = une entrée
`{index, style, texte}`, texte IDENTIQUE au fichier source (aucune réécriture,
aucun résumé, aucun appel LLM à cette étape). Les tableaux éventuels du
document (ex. répartition du capital) sont extraits séparément dans le même
fichier (`tables`), cellule par cellule.

POURQUOI un script séparé de `extract_pacte_fondateurs_clauses.py` (Phase 1,
découpage en clauses) : la Phase 0 n'a aucune règle métier, juste une lecture
mécanique du `.docx` via `python-docx` — la garder isolée permet de vérifier
le résultat contre le fichier Word original AVANT tout découpage, plutôt que
de devoir faire confiance à un script qui lit ET interprète en une seule
passe (c'est cette confiance aveugle, sur la première extraction de ce
service, qui a produit un JSON final incohérent avec le texte source réel —
voir la discussion du 20/09/2026 avec l'utilisateur).

N'écrit RIEN en base : ce script ne dépend pas de `backend.shared.database`,
contrairement aux autres scripts de `knowledge_base/scripts/` (ex.
`collect_legal_texts_formes_juridiques.py`) qui font un fetch + upsert direct.
Ici la source est un fichier local déjà en main, pas une API à interroger —
la seule chose à garantir est la fidélité de la lecture, vérifiable par
l'utilisateur avant que la Phase 1 ne s'en serve.
"""
import argparse
import json
from pathlib import Path

import docx

_OUTPUT_DIR = Path(__file__).parent.parent / "data" / "pacte_fondateurs"


def extract_raw(source_path: str, label: str) -> Path:
    document = docx.Document(source_path)

    paragraphes = []
    for index, paragraph in enumerate(document.paragraphs):
        texte = paragraph.text
        if texte.strip() == "":
            # Paragraphes vides = mise en forme Word (espacement entre sections),
            # aucun contenu à préserver — mais l'index original n'est PAS
            # renuméroté après filtrage : il reste celui de python-docx, pour
            # pouvoir recroiser avec le .docx si besoin de vérification.
            continue
        style = paragraph.style.name if paragraph.style else None
        paragraphes.append({"index": index, "style": style, "texte": texte})

    tables = []
    for table_index, table in enumerate(document.tables):
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        tables.append({"table_index": table_index, "rows": rows})

    _OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    outpath = _OUTPUT_DIR / f"raw_{label}.json"
    with open(outpath, "w", encoding="utf-8") as f:
        json.dump(
            {"source_file": label, "paragraphes": paragraphes, "tables": tables},
            f,
            ensure_ascii=False,
            indent=2,
        )
    return outpath


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", required=True, help="Chemin vers le fichier .docx source.")
    parser.add_argument(
        "--label",
        required=True,
        help="Nom court identifiant la source (ex: modele_principal, notice_explicative) — "
        "détermine le nom du fichier de sortie raw_<label>.json.",
    )
    args = parser.parse_args()

    outpath = extract_raw(args.source, args.label)
    with open(outpath, encoding="utf-8") as f:
        data = json.load(f)
    print(
        f"OK — {len(data['paragraphes'])} paragraphe(s) non vide(s), "
        f"{len(data['tables'])} tableau(x) -> {outpath}"
    )


if __name__ == "__main__":
    main()
