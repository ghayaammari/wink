"""Rendu d'un DeckContentModel (agents/pitch_deck_agent/deck_schema.py) en

fichier PPTX brande (service 6.6.2). Port quasi-integral du renderer du
prototype startapini (agents/pitch-deck-agent/app/renderers/pptx_renderer.py),
simplifie pour une seule langue (voir deck_schema.py) et sans le chemin de
rendu externe via MCP/presentation-engine (non porte, decision du 21/09/2026 --
"reste sobre", ce service Node n'existe pas dans ce repo).

Aucun appel LLM ici : purement deterministe (python-pptx), rendu a la volee a
chaque telechargement (voir pitch_router.py) plutot que stocke, meme principe
que compliance_report_service.py pour les rapports PDF RGPD/AI Act. Une
fonction par diapositive, les 11 diapositives de DeckSlides.
"""

from __future__ import annotations

import io
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from agents.pitch_deck_agent.deck_schema import DeckContentModel, MarketStat, TractionMetric, TractionMilestone
from backend.entrepreneur.services.pitch_deck_assets import fetch_logo_image, get_initials
from backend.entrepreneur.services.pitch_deck_theme import SLIDE_HEIGHT_IN, SLIDE_WIDTH_IN, ResolvedTheme, resolve_theme

BLANK_LAYOUT_INDEX = 6
CHECK_GLYPH = "✓"
ARROW_GLYPH = "→"


# --------------------------------------------------------------------------
# Low-level drawing helpers
# --------------------------------------------------------------------------


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.lstrip("#").upper())


def _new_slide(prs: Presentation, theme: ResolvedTheme):
    slide = prs.slides.add_slide(prs.slide_layouts[BLANK_LAYOUT_INDEX])
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = _rgb(theme.background)
    return slide


def _set_font(font, *, name: str, size: float, color: str, bold: bool = False, italic: bool = False) -> None:
    font.name = name
    font.size = Pt(size)
    font.bold = bold
    font.italic = italic
    font.color.rgb = _rgb(color)


def _add_textbox(
    slide,
    text: str,
    left: float,
    top: float,
    width: float,
    height: float,
    *,
    font: str,
    size: float,
    color: str,
    bold: bool = False,
    italic: bool = False,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    _set_font(run.font, name=font, size=size, color=color, bold=bold, italic=italic)
    return box


def _add_shape(slide, shape_type, left: float, top: float, width: float, height: float, *, fill_color: str | None):
    shape = slide.shapes.add_shape(shape_type, Inches(left), Inches(top), Inches(width), Inches(height))
    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(fill_color)
    else:
        shape.fill.background()
    shape.line.fill.background()
    return shape


def _shape_text(
    shape,
    text: str,
    *,
    font: str,
    size: float,
    color: str,
    bold: bool = False,
    align=PP_ALIGN.CENTER,
) -> None:
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    _set_font(run.font, name=font, size=size, color=color, bold=bold)


def _add_soft_shadow(shape) -> None:
    """Subtle outer drop shadow -- python-pptx's high-level API only exposes
    shadow.inherit, so this reaches into the raw XML effect list."""
    spPr = shape._element.spPr
    effect_lst = spPr.makeelement(qn("a:effectLst"), {})
    outer_shadow = spPr.makeelement(
        qn("a:outerShdw"),
        {"blurRad": "90000", "dist": "30000", "dir": "5400000", "rotWithShape": "0"},
    )
    color = spPr.makeelement(qn("a:srgbClr"), {"val": "000000"})
    alpha = spPr.makeelement(qn("a:alpha"), {"val": "35000"})
    color.append(alpha)
    outer_shadow.append(color)
    effect_lst.append(outer_shadow)
    spPr.append(effect_lst)


def _crop_picture_to_ellipse(picture) -> None:
    """Masks a rectangular picture into a circle by swapping its shape
    geometry to an ellipse -- python-pptx has no picture-fill/mask API, this
    is the standard low-level workaround (PowerPoint honors prstGeom on
    pictures the same way it does on autoshapes)."""
    geom = picture._element.spPr.find(qn("a:prstGeom"))
    if geom is not None:
        geom.set("prst", "ellipse")


def _set_bullet(paragraph, *, color: str, char: str = "•") -> None:
    """Native bullet glyph via raw XML -- python-pptx's object model doesn't
    expose bullet formatting at all."""
    pPr = paragraph._p.get_or_add_pPr()
    bu_clr = pPr.makeelement(qn("a:buClr"), {})
    bu_clr.append(pPr.makeelement(qn("a:srgbClr"), {"val": color.lstrip("#").upper()}))
    bu_font = pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"})
    bu_char = pPr.makeelement(qn("a:buChar"), {"char": char})
    pPr.append(bu_clr)
    pPr.append(bu_font)
    pPr.append(bu_char)
    pPr.set("marL", str(Emu(Inches(0.28))))
    pPr.set("indent", str(-Emu(Inches(0.28))))


def _circle_badge(slide, *, cx: float, cy: float, diameter: float, fill_color: str):
    left, top = cx - diameter / 2, cy - diameter / 2
    return _add_shape(slide, MSO_SHAPE.OVAL, left, top, diameter, diameter, fill_color=fill_color)


# --------------------------------------------------------------------------
# Persistent logo badges (identical position on every slide)
# --------------------------------------------------------------------------

_BADGE_Y = SLIDE_HEIGHT_IN - 0.65
_BADGE_SIZE = 0.42
_WINK_BADGE_X = 0.9
_STARTUP_BADGE_X = SLIDE_WIDTH_IN - 0.9 - _BADGE_SIZE


def _add_footer_badges(slide, model: DeckContentModel, theme: ResolvedTheme) -> None:
    # Marque de plateforme Wink : en bas a gauche, discrete.
    mark = _add_shape(
        slide, MSO_SHAPE.ROUNDED_RECTANGLE, _WINK_BADGE_X, _BADGE_Y, _BADGE_SIZE, _BADGE_SIZE, fill_color=theme.icon_circle_bg
    )
    _shape_text(mark, "W", font=theme.title_font, size=16, color=theme.text_muted, bold=True)
    _add_textbox(
        slide,
        "Wink",
        _WINK_BADGE_X + _BADGE_SIZE + 0.12,
        _BADGE_Y,
        1.4,
        _BADGE_SIZE,
        font=theme.body_font,
        size=10,
        color=theme.text_muted,
        anchor=MSO_ANCHOR.MIDDLE,
    )

    # Startup mark: bottom-right, prominent, name to its left.
    startup_name = model.slides.title.startup_name
    _add_textbox(
        slide,
        startup_name,
        _STARTUP_BADGE_X - 2.6,
        _BADGE_Y,
        2.4,
        _BADGE_SIZE,
        font=theme.body_font,
        size=10,
        color=theme.text_muted,
        align=PP_ALIGN.RIGHT,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    badge_cx = _STARTUP_BADGE_X + _BADGE_SIZE / 2
    badge_cy = _BADGE_Y + _BADGE_SIZE / 2

    logo_bytes = fetch_logo_image(model.brand.logo_url, startup_id=model.project_id)
    if logo_bytes is not None:
        picture = slide.shapes.add_picture(
            io.BytesIO(logo_bytes),
            Inches(_STARTUP_BADGE_X),
            Inches(_BADGE_Y),
            Inches(_BADGE_SIZE),
            Inches(_BADGE_SIZE),
        )
        _crop_picture_to_ellipse(picture)
    else:
        badge = _circle_badge(slide, cx=badge_cx, cy=badge_cy, diameter=_BADGE_SIZE, fill_color=theme.accent_color)
        _shape_text(badge, get_initials(startup_name), font=theme.title_font, size=14, color=theme.text_primary, bold=True)


# --------------------------------------------------------------------------
# Slide header
# --------------------------------------------------------------------------


def _slide_header(slide, title: str, theme: ResolvedTheme) -> float:
    _add_textbox(slide, title, 0.9, 0.55, SLIDE_WIDTH_IN - 1.8, 0.9, font=theme.title_font, size=32, color=theme.text_primary, bold=True)
    return 1.55  # y-coordinate content should start below the header


# --------------------------------------------------------------------------
# Icon row (Problem: numbered: Business model: checkmark)
# --------------------------------------------------------------------------


def _render_icon_rows(slide, theme: ResolvedTheme, bullets: list[str], *, glyph_for: "callable", content_top: float) -> None:
    row_height = 1.15
    circle_d = 0.55
    left = 0.9
    text_left = left + circle_d + 0.35
    text_width = SLIDE_WIDTH_IN - text_left - 0.9

    for i, bullet in enumerate(bullets):
        row_top = content_top + i * row_height
        circle = _circle_badge(slide, cx=left + circle_d / 2, cy=row_top + row_height / 2, diameter=circle_d, fill_color=theme.icon_circle_bg)
        _shape_text(circle, glyph_for(i), font=theme.title_font, size=18, color=theme.accent_color, bold=True)
        _add_textbox(
            slide,
            bullet,
            text_left,
            row_top,
            text_width,
            row_height,
            font=theme.body_font,
            size=16,
            color=theme.text_primary,
            anchor=MSO_ANCHOR.MIDDLE,
        )


def _render_problem(prs: Presentation, model: DeckContentModel, theme: ResolvedTheme) -> None:
    slide = _new_slide(prs, theme)
    problem = model.slides.problem
    content_top = _slide_header(slide, problem.title, theme)
    bullets = problem.bullets
    _render_icon_rows(slide, theme, bullets, glyph_for=lambda i: str(i + 1), content_top=content_top)
    _add_footer_badges(slide, model, theme)


def _render_business_model(prs: Presentation, model: DeckContentModel, theme: ResolvedTheme) -> None:
    slide = _new_slide(prs, theme)
    bm = model.slides.business_model
    content_top = _slide_header(slide, bm.title, theme)
    bullets = bm.bullets
    _render_icon_rows(slide, theme, bullets, glyph_for=lambda i: CHECK_GLYPH, content_top=content_top)
    _add_footer_badges(slide, model, theme)


def _render_competition(prs: Presentation, model: DeckContentModel, theme: ResolvedTheme) -> None:
    slide = _new_slide(prs, theme)
    competition = model.slides.competition
    content_top = _slide_header(slide, competition.title, theme)
    bullets = competition.bullets
    _render_icon_rows(slide, theme, bullets, glyph_for=lambda i: CHECK_GLYPH, content_top=content_top)
    _add_footer_badges(slide, model, theme)


def _render_financials(prs: Presentation, model: DeckContentModel, theme: ResolvedTheme) -> None:
    slide = _new_slide(prs, theme)
    financials = model.slides.financials
    content_top = _slide_header(slide, financials.title, theme)
    bullets = financials.bullets
    _render_icon_rows(slide, theme, bullets, glyph_for=lambda i: CHECK_GLYPH, content_top=content_top)
    _add_footer_badges(slide, model, theme)


def _render_roadmap(prs: Presentation, model: DeckContentModel, theme: ResolvedTheme) -> None:
    slide = _new_slide(prs, theme)
    roadmap = model.slides.roadmap
    content_top = _slide_header(slide, roadmap.title, theme)
    bullets = roadmap.bullets
    _render_icon_rows(slide, theme, bullets, glyph_for=lambda i: CHECK_GLYPH, content_top=content_top)
    _add_footer_badges(slide, model, theme)


# --------------------------------------------------------------------------
# Icon grid (Solution: numbered, CTA: arrow)
# --------------------------------------------------------------------------


def _render_icon_grid(slide, theme: ResolvedTheme, bullets: list[str], *, glyph_for, content_top: float, content_bottom: float) -> None:
    columns, rows = 2, 2
    gutter = 0.35
    left, right_margin = 0.9, 0.9
    grid_width = SLIDE_WIDTH_IN - left - right_margin
    grid_height = content_bottom - content_top
    card_w = (grid_width - gutter * (columns - 1)) / columns
    card_h = (grid_height - gutter * (rows - 1)) / rows

    for i, bullet in enumerate(bullets[: columns * rows]):
        col, row = i % columns, i // columns
        card_left = left + col * (card_w + gutter)
        card_top = content_top + row * (card_h + gutter)
        card = _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, card_left, card_top, card_w, card_h, fill_color=theme.card_background)
        card.adjustments[0] = 0.06

        circle_d = 0.5
        circle = _circle_badge(slide, cx=card_left + 0.5, cy=card_top + 0.5, diameter=circle_d, fill_color=theme.icon_circle_bg)
        _shape_text(circle, glyph_for(i), font=theme.title_font, size=16, color=theme.accent_color, bold=True)

        _add_textbox(
            slide,
            bullet,
            card_left + 0.25,
            card_top + 0.9,
            card_w - 0.5,
            card_h - 1.1,
            font=theme.body_font,
            size=14,
            color=theme.text_primary,
        )


def _render_solution(prs: Presentation, model: DeckContentModel, theme: ResolvedTheme) -> None:
    slide = _new_slide(prs, theme)
    solution = model.slides.solution
    content_top = _slide_header(slide, solution.title, theme)
    _render_icon_grid(slide, theme, solution.bullets, glyph_for=lambda i: str(i + 1), content_top=content_top, content_bottom=SLIDE_HEIGHT_IN - 1.1)
    _add_footer_badges(slide, model, theme)


def _render_cta(prs: Presentation, model: DeckContentModel, theme: ResolvedTheme) -> None:
    slide = _new_slide(prs, theme)
    cta = model.slides.cta
    content_top = _slide_header(slide, cta.title, theme)

    ask_card_height = 1.0
    ask_card = _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.9, content_top, SLIDE_WIDTH_IN - 1.8, ask_card_height, fill_color=theme.card_background)
    ask_card.adjustments[0] = 0.08
    _shape_text(ask_card, cta.ask_summary, font=theme.body_font, size=18, color=theme.accent_color, bold=True)
    ask_card.text_frame.paragraphs[0].runs[0].font.italic = True

    grid_top = content_top + ask_card_height + 0.3
    _render_icon_grid(slide, theme, cta.bullets, glyph_for=lambda i: ARROW_GLYPH, content_top=grid_top, content_bottom=SLIDE_HEIGHT_IN - 1.1)
    _add_footer_badges(slide, model, theme)


# --------------------------------------------------------------------------
# Market: bullets + stat callout + growth badge
# --------------------------------------------------------------------------


def _render_stat_callout(slide, theme: ResolvedTheme, stat: MarketStat, *, left: float, top: float, width: float) -> float:
    _add_textbox(slide, stat.value, left, top, width, 1.0, font=theme.title_font, size=54, color=theme.accent_color, bold=True, align=PP_ALIGN.CENTER)
    _add_textbox(slide, stat.label, left, top + 0.95, width, 0.5, font=theme.body_font, size=13, color=theme.text_muted, align=PP_ALIGN.CENTER)
    return top + 1.55


def _render_mini_bars(slide, theme: ResolvedTheme, *, left: float, top: float, width: float) -> float:
    bar_count, bar_w, gap = 4, 0.35, 0.15
    max_h = 0.8
    total_w = bar_count * bar_w + (bar_count - 1) * gap
    start_x = left + (width - total_w) / 2
    for i in range(bar_count):
        h = max_h * (i + 1) / bar_count
        color = theme.accent_color if i == bar_count - 1 else theme.icon_circle_bg
        _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, start_x + i * (bar_w + gap), top + (max_h - h), bar_w, h, fill_color=color)
    return top + max_h + 0.2


def _render_growth_badge(slide, theme: ResolvedTheme, highlight: MarketStat, *, left: float, top: float, width: float) -> None:
    badge_w, badge_h = min(width, 2.6), 0.55
    badge_left = left + (width - badge_w) / 2
    chip = _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, badge_left, top, badge_w, badge_h, fill_color=theme.icon_circle_bg)
    chip.adjustments[0] = 0.5
    _shape_text(chip, f"{highlight.value}  {highlight.label}", font=theme.body_font, size=12, color=theme.accent_color, bold=True)


def _render_market(prs: Presentation, model: DeckContentModel, theme: ResolvedTheme) -> None:
    slide = _new_slide(prs, theme)
    market = model.slides.market
    content_top = _slide_header(slide, market.title, theme)

    left_width = 6.4
    box = slide.shapes.add_textbox(Inches(0.9), Inches(content_top), Inches(left_width), Inches(SLIDE_HEIGHT_IN - content_top - 1.1))
    tf = box.text_frame
    tf.word_wrap = True
    for i, bullet in enumerate(market.bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = bullet
        p.space_after = Pt(14)
        _set_font(p.font, name=theme.body_font, size=16, color=theme.text_primary)
        _set_bullet(p, color=theme.accent_color)

    right_left = 0.9 + left_width + 0.4
    right_width = SLIDE_WIDTH_IN - right_left - 0.9
    y = content_top + 0.2
    if market.stat:
        y = _render_stat_callout(slide, theme, market.stat, left=right_left, top=y, width=right_width)
        y = _render_mini_bars(slide, theme, left=right_left, top=y, width=right_width)
    if market.highlight:
        _render_growth_badge(slide, theme, market.highlight, left=right_left, top=y, width=right_width)

    _add_footer_badges(slide, model, theme)


# --------------------------------------------------------------------------
# Traction: stat row + timeline
# --------------------------------------------------------------------------


def _render_traction_stats(slide, theme: ResolvedTheme, metrics: list[TractionMetric], *, top: float) -> None:
    if not metrics:
        return
    left, right_margin = 0.9, 0.9
    width = SLIDE_WIDTH_IN - left - right_margin
    col_w = width / len(metrics)

    for i, metric in enumerate(metrics):
        cx = left + col_w * i + col_w / 2
        ring = _add_shape(slide, MSO_SHAPE.OVAL, cx - 0.22, top, 0.44, 0.44, fill_color=None)
        ring.line.fill.solid()
        ring.line.fill.fore_color.rgb = _rgb(theme.accent_color)
        ring.line.width = Pt(1.5)
        _add_textbox(slide, metric.value, cx - col_w / 2, top + 0.55, col_w, 0.8, font=theme.title_font, size=int(min(48, 44)), color=theme.accent_color, bold=True, align=PP_ALIGN.CENTER)
        _add_textbox(slide, metric.label, cx - col_w / 2, top + 1.35, col_w, 0.5, font=theme.body_font, size=13, color=theme.text_muted, align=PP_ALIGN.CENTER)


def _render_timeline(slide, theme: ResolvedTheme, milestones: list[TractionMilestone], *, top: float) -> None:
    if not milestones:
        return
    left, right_margin = 1.2, 1.2
    width = SLIDE_WIDTH_IN - left - right_margin

    line = _add_shape(slide, MSO_SHAPE.RECTANGLE, left, top, width, 0.03, fill_color=theme.icon_circle_bg)

    count = len(milestones)
    for i, milestone in enumerate(milestones):
        x = left if count == 1 else left + width * i / (count - 1)
        dot_d = 0.16
        _circle_badge(slide, cx=x, cy=top + 0.015, diameter=dot_d, fill_color=theme.accent_color)
        _add_textbox(slide, milestone.date, x - 0.9, top - 0.55, 1.8, 0.4, font=theme.body_font, size=13, color=theme.accent_color, bold=True, align=PP_ALIGN.CENTER)
        _add_textbox(slide, milestone.label, x - 0.9, top + 0.2, 1.8, 0.6, font=theme.body_font, size=12, color=theme.text_primary, align=PP_ALIGN.CENTER)


def _render_traction(prs: Presentation, model: DeckContentModel, theme: ResolvedTheme) -> None:
    slide = _new_slide(prs, theme)
    traction = model.slides.traction
    content_top = _slide_header(slide, traction.title, theme)

    _render_traction_stats(slide, theme, traction.metrics, top=content_top + 0.1)
    _render_timeline(slide, theme, traction.milestones, top=content_top + 3.1)
    _add_footer_badges(slide, model, theme)


# --------------------------------------------------------------------------
# Team: cards with avatar (photo or initials)
# --------------------------------------------------------------------------


def _render_team(prs: Presentation, model: DeckContentModel, theme: ResolvedTheme) -> None:
    slide = _new_slide(prs, theme)
    team = model.slides.team
    content_top = _slide_header(slide, team.title, theme)

    members = team.members
    columns = min(len(members), 4)
    rows = -(-len(members) // columns)  # ceil
    gutter = 0.3
    left, right_margin = 0.9, 0.9
    grid_width = SLIDE_WIDTH_IN - left - right_margin
    content_bottom = SLIDE_HEIGHT_IN - 1.1
    card_w = (grid_width - gutter * (columns - 1)) / columns
    card_h = min(3.2, (content_bottom - content_top - gutter * (rows - 1)) / rows)

    for i, member in enumerate(members):
        col, row = i % columns, i // columns
        card_left = left + col * (card_w + gutter)
        card_top = content_top + row * (card_h + gutter)
        card = _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, card_left, card_top, card_w, card_h, fill_color=theme.card_background)
        card.adjustments[0] = 0.08
        _add_soft_shadow(card)

        avatar_d = 0.95
        avatar_cx = card_left + card_w / 2
        avatar_cy = card_top + 0.65
        avatar = _circle_badge(slide, cx=avatar_cx, cy=avatar_cy, diameter=avatar_d, fill_color=theme.accent_color)

        photo_bytes = fetch_logo_image(member.photo_url, startup_id=f"{model.project_id}:{member.name}")
        if photo_bytes is not None:
            picture = slide.shapes.add_picture(
                io.BytesIO(photo_bytes), Inches(avatar_cx - avatar_d / 2), Inches(avatar_cy - avatar_d / 2), Inches(avatar_d), Inches(avatar_d)
            )
            _crop_picture_to_ellipse(picture)
        else:
            _shape_text(avatar, get_initials(member.name), font=theme.title_font, size=20, color=theme.text_primary, bold=True)

        _add_textbox(slide, member.name, card_left + 0.1, card_top + 1.25, card_w - 0.2, 0.4, font=theme.body_font, size=15, color=theme.text_primary, bold=True, align=PP_ALIGN.CENTER)
        _add_textbox(slide, member.role, card_left + 0.1, card_top + 1.65, card_w - 0.2, 0.4, font=theme.body_font, size=12, color=theme.text_muted, align=PP_ALIGN.CENTER)

    _add_footer_badges(slide, model, theme)


# --------------------------------------------------------------------------
# Title
# --------------------------------------------------------------------------


def _render_title(prs: Presentation, model: DeckContentModel, theme: ResolvedTheme) -> None:
    slide = _new_slide(prs, theme)
    _add_textbox(
        slide, "REGBRIDGE STARTUP SPOTLIGHT", 1.0, 2.2, SLIDE_WIDTH_IN - 2.0, 0.5,
        font=theme.body_font, size=12, color=theme.text_muted, align=PP_ALIGN.CENTER,
    )
    _add_textbox(
        slide, model.slides.title.startup_name, 1.0, 2.9, SLIDE_WIDTH_IN - 2.0, 1.2,
        font=theme.title_font, size=48, color=theme.text_primary, bold=True, align=PP_ALIGN.CENTER,
    )
    _add_textbox(
        slide, model.slides.title.one_liner, 1.5, 4.15, SLIDE_WIDTH_IN - 3.0, 1.0,
        font=theme.body_font, size=20, color=theme.accent_color, italic=True, align=PP_ALIGN.CENTER,
    )
    _add_footer_badges(slide, model, theme)


# --------------------------------------------------------------------------
# Entry point
# --------------------------------------------------------------------------


def render_pptx(model: DeckContentModel) -> bytes:
    theme = resolve_theme(model.brand.accent_color)
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_WIDTH_IN)
    prs.slide_height = Inches(SLIDE_HEIGHT_IN)

    _render_title(prs, model, theme)
    _render_problem(prs, model, theme)
    _render_solution(prs, model, theme)
    _render_market(prs, model, theme)
    _render_business_model(prs, model, theme)
    _render_competition(prs, model, theme)
    _render_traction(prs, model, theme)
    _render_team(prs, model, theme)
    _render_financials(prs, model, theme)
    _render_roadmap(prs, model, theme)
    _render_cta(prs, model, theme)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()

